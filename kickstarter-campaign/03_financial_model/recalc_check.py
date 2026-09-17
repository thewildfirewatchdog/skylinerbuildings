import formulas, sys
xl = formulas.ExcelModel().loads('WW_Kickstarter_Financial_Model_v1.xlsx').finish()
sol = xl.calculate()
def v(sheet, cell):
    k=f"'[WW_Kickstarter_Financial_Model_v1.xlsx]{sheet.upper()}'!{cell}"
    r = sol.get(k)
    return None if r is None else r.value[0][0]
errs=[k for k,r in sol.items() if hasattr(r,'value') and isinstance(r.value[0][0],str) and str(r.value[0][0]).startswith('#')]
print('formula errors:', errs[:10])
from openpyxl import load_workbook
wb=load_workbook('WW_Kickstarter_Financial_Model_v1.xlsx')
for sh in ['Unit_Cost','Rewards','Scenarios','Funding_Goal','Cash_Timing','Stress_Tests','Capacity']:
    ws=wb[sh]; print('=== ',sh)
    for row in ws.iter_rows(min_row=2):
        label=row[0].value
        if not label: continue
        vals=[]
        for c in row[1:7]:
            if c.value is None: continue
            val = v(sh,c.coordinate) if isinstance(c.value,str) and c.value.startswith('=') else c.value
            if isinstance(val,float): val=round(val,1)
            vals.append(val)
        print(f"{str(label)[:62]:62}", vals[:6])
